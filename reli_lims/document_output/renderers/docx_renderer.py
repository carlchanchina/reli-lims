"""Expand designated template rows without rebuilding their cell formatting."""
from copy import deepcopy
from io import BytesIO
from docx import Document


def replace(paragraph, values):
    # Templates keep placeholders in a single run so fonts and paragraph styles survive.
    for run in paragraph.runs:
        for key, value in values.items():
            run.text = run.text.replace('{{' + key + '}}', str(value if value is not None else ''))


def expand(table, marker, rows, keys):
    for row in list(table.rows):
        if marker not in ' '.join(cell.text for cell in row.cells):
            continue
        for values in rows:
            element = deepcopy(row._tr)
            row._tr.addprevious(element)
            from docx.table import _Row
            copied = _Row(element, table)
            mapping = dict(zip(keys, values))
            for cell in copied.cells:
                for paragraph in cell.paragraphs:
                    replace(paragraph, mapping)
        table._tbl.remove(row._tr)


def render(spec, context):
    doc = Document(spec.directory / f'{spec.key}.docx')
    scalars = {key: context.get(key, '') for key in ('title', 'name', 'notice', 'terms')}
    for p in doc.paragraphs:
        replace(p, scalars)
    for table in doc.tables:
        expand(table, '{{field_label}}', context['fields'], ['field_label', 'field_value'])
        expand(table, '{{row.0}}', context['rows'], [f'row.{i}' for i in range(len(spec.columns))])
    for section in context.get('sections', []):
        doc.add_heading(section['title'], level=2)
        table = doc.add_table(rows=1, cols=len(section['columns']))
        table.style = 'Table Grid'
        for c, value in zip(table.rows[0].cells, section['columns']):
            c.text = str(value)
        for row in section['rows']:
            for c, value in zip(table.add_row().cells, row):
                c.text = str(value if value is not None else '')
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
