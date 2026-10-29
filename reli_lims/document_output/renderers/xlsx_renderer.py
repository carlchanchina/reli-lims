"""Fill existing workbooks; preserve styles, merged footer, dimensions and print layout."""
from copy import copy
from io import BytesIO
from math import ceil
import re
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import range_boundaries, get_column_letter


def write(cell, value):
    # Split before assigning: openpyxl silently truncates at Excel's cell limit.
    if isinstance(value, str) and len(value) > 32767:
        value = continued_value(cell, value)
    cell.value = value if value is not None else ''
    # Customer text must never be interpreted as Excel formulas.
    if isinstance(value, str):
        cell.data_type = 's'


def expand_rows(sheet, start, reserved, count):
    extra = max(0, count - reserved)
    if not extra:
        return
    pivot = start + reserved
    merges = [str(r) for r in sheet.merged_cells.ranges]
    dimensions = {i: copy(dim) for i, dim in sheet.row_dimensions.items()}
    for area in merges:
        sheet.unmerge_cells(area)
    sheet.insert_rows(pivot, extra)
    for i, dimension in sorted(dimensions.items(), reverse=True):
        if i >= pivot:
            sheet.row_dimensions[i + extra] = dimension
            sheet.row_dimensions[i + extra].index = i + extra
    for offset in range(extra):
        dest = pivot + offset
        sheet.row_dimensions[dest] = copy(dimensions.get(pivot - 1, sheet.row_dimensions[start]))
        sheet.row_dimensions[dest].index = dest
        for col in range(1, sheet.max_column + 1):
            source, target = sheet.cell(pivot - 1, col), sheet.cell(dest, col)
            target._style = copy(source._style)
            if source.data_type == 'f':
                target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
    for area in merges:
        left, top, right, bottom = range_boundaries(area)
        if top >= pivot:
            top += extra
        if bottom >= pivot:
            bottom += extra
        sheet.merge_cells(start_row=top, start_column=left, end_row=bottom, end_column=right)
    # The packaged templates have one contiguous A4 print area.
    sheet.print_area = f'A1:F{sheet.max_row}'


def append_wide(sheet, values, columns):
    sheet.append(values)
    for column, value in enumerate(values, 1):
        write(sheet.cell(sheet.max_row, column), value)
    if len(values) in (1, 2):
        sheet.merge_cells(start_row=sheet.max_row, end_row=sheet.max_row,
                          start_column=len(values), end_column=columns)


def line_count(value, width):
    return sum(max(1, ceil(len(line) * 2 / max(width, 1))) for line in str(value).split('\n'))


def continuation(workbook, source, original, number):
    if '续页' not in workbook.sheetnames:
        appendix = workbook.create_sheet('续页')
        appendix.page_setup.paperSize = appendix.PAPERSIZE_A4
        appendix.page_setup.fitToWidth = 1
        appendix.page_setup.fitToHeight = 0
        appendix.sheet_properties.pageSetUpPr.fitToPage = True
        appendix.print_title_rows = '1:1'
        for col in 'ABCDEF':
            appendix.column_dimensions[col].width = 18
        append_wide(appendix, ['文档续页'], 6)
    appendix = workbook['续页']
    append_wide(appendix, [f'第 {number} 节 · {source}'], 6)
    # Each merged row remains well below Excel's row-height limit, including explicit newlines.
    chunk = ''
    for character in original:
        if chunk and (len(chunk) >= 250 or line_count(chunk + character, 108) > 12):
            append_wide(appendix, [chunk], 6)
            chunk = ''
        chunk += character
    if chunk:
        append_wide(appendix, [chunk], 6)
    for row in appendix:
        cell = row[0]
        cell.font = Font(name='宋体', size=11)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        appendix.row_dimensions[cell.row].height = max(25, line_count(cell.value, 108) * 17)
    appendix.print_area = f'A1:F{appendix.max_row}'


def continued_value(cell, original, prefix_length=50):
    workbook = cell.parent.parent
    number = getattr(workbook, '_reli_continuation_count', 0) + 1
    workbook._reli_continuation_count = number
    continuation(workbook, f'{cell.parent.title}!{cell.coordinate}', original, number)
    return original[:prefix_length] + f'…\n（全文见续页第 {number} 节）'


def layout(workbook, sheet):
    for row in sheet:
        height = sheet.row_dimensions[row[0].row].height or 16
        for cell in row:
            if cell.value is None:
                continue
            width = sheet.column_dimensions[cell.column_letter].width or 13
            for area in sheet.merged_cells.ranges:
                if cell.coordinate in area:
                    width = sum(sheet.column_dimensions[get_column_letter(c)].width or 13
                                for c in range(area.min_col, area.max_col + 1))
                    break
            lines = line_count(cell.value, width)
            if lines > 20:
                prefix_length = max(20, min(150, int(width / 2) * 5))
                write(cell, continued_value(cell, str(cell.value), prefix_length))
                lines = line_count(cell.value, width)
            if lines > 1:
                cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical='top', wrap_text=True)
                height = max(height, lines * 17)
        sheet.row_dimensions[row[0].row].height = height


def render(spec, context):
    workbook = load_workbook(spec.directory / f'{spec.key}.xlsx')
    sheet = workbook.active
    expand_rows(sheet, spec.row_start, spec.reserved_rows, len(context['rows']))
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str) and '{{' in cell.value:
                def replacement(match):
                    value = context.get(match.group(1))
                    return str(value) if isinstance(value, (str, int, float)) else ''
                value = re.sub(r'\{\{([^}]+)\}\}', replacement, cell.value)
                write(cell, value)
    for i in range(max(spec.reserved_rows, len(context['rows']))):
        values = context['rows'][i] if i < len(context['rows']) else []
        for col in range(1, len(spec.columns) + 1):
            value = values[col - 1] if col <= len(values) else ''
            # Numeric values stay numeric for accounting and editing.
            if spec.key == 'quotation' and col in (4, 5, 6) and value != '':
                from decimal import Decimal, InvalidOperation
                try:
                    value = float(Decimal(str(value)))
                except InvalidOperation:
                    pass
            write(sheet.cell(spec.row_start + i, col), value)
    if spec.key == 'quotation':
        total_row = spec.row_start + max(spec.reserved_rows, len(context['rows']))
        write(sheet.cell(total_row, 6), context.get('total', ''))
    else:
        # Metadata and customer-facing sections follow the detail table.
        sheet.append([])
        for label, value in context['fields']:
            append_wide(sheet, [label, value], len(spec.columns))
        for section in context.get('sections', []):
            append_wide(sheet, [section['title']], len(spec.columns))
            append_wide(sheet, section['columns'], len(spec.columns))
            for row in section['rows']:
                append_wide(sheet, [str(v) if v is not None else '' for v in row], len(spec.columns))
        append_wide(sheet, ['编制：____________　审核：____________　日期：____________'], len(spec.columns))
        prototype = sheet.cell(spec.row_start, 1)
        for row in sheet:
            for cell in row:
                if cell.row > spec.row_start + max(spec.reserved_rows, len(context['rows'])) + 2:
                    cell._style = copy(prototype._style)
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.data_type == 'f':
                    write(cell, cell.value)
        sheet.print_area = f'A1:{get_column_letter(len(spec.columns))}{sheet.max_row}'
    layout(workbook, sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
