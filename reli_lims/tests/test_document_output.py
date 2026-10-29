"""Run on a disposable Frappe site: real permissions, native PDF and private files."""
from io import BytesIO
import uuid

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import today
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader

from reli_lims.document_output.service import generate_document


class TestDocumentOutput(IntegrationTestCase):
    def setUp(self):
        super().setUp()
        frappe.set_user('Administrator')
        self.addCleanup(frappe.set_user, 'Administrator')
        token = uuid.uuid4().hex[:8]
        self.company = frappe.get_doc(dict(doctype='Company', company_name=f'Output {token}',
            abbr=f'O{token}', default_currency='USD', country='United States', chart_of_accounts='Standard')).insert()
        self.customer = frappe.get_doc(dict(doctype='Customer', customer_name=f'文档示例 {token}', customer_type='Company')).insert()
        method = frappe.get_doc(dict(doctype='Test Method', method_name=f'低温 {token}')).insert()
        standard = frappe.get_doc(dict(doctype='Test Standard', standard_code=f'A-{token}', title='示例标准', version='2026')).insert()
        self.request = frappe.get_doc(dict(doctype='Test Request', title='文档输出验收', customer=self.customer.name,
            company=self.company.name, request_date=today(), delivery_date=today(),
            samples=[dict(sample_code='S1', product_name='示例样品', quantity=1)],
            items=[dict(item_code='T1', test_method=method.name, standard=standard.name)],
            sample_assignments=[dict(item_code='T1', sample_code='S1')])).insert()
        self.project = frappe.get_doc(dict(doctype='Project', project_name=f'Output {token}',
            customer=self.customer.name, company=self.company.name)).insert()
        sample = frappe.get_doc(dict(doctype='Sample', project=self.project.name, received_date=today(), quantity=1)).insert()
        self.program = frappe.get_doc(dict(doctype='Test Program', project=self.project.name,
            samples=[dict(sample=sample.name)],
            items=[dict(item_code='T1', test_method=method.name, standard=standard.name)],
            sample_assignments=[dict(item_code='T1', sample=sample.name)])).insert()
        self.report = frappe.get_doc(dict(doctype='Test Report', project=self.project.name,
            test_program=self.program.name, overall_conclusion='仅测试数据')).insert()
        item = frappe.get_doc(dict(doctype='Item', item_code=f'OUTPUT-{token}', item_name='示例检测服务',
            item_group='All Item Groups', stock_uom='Nos', is_stock_item=0)).insert()
        self.quotation = frappe.get_doc(dict(doctype='Quotation', quotation_to='Customer', party_name=self.customer.name,
            company=self.company.name, transaction_date=today(), currency='USD', selling_price_list='Standard Selling',
            items=[dict(item_code=item.name, qty=2, rate=100)])).insert()

    def test_all_nine_formats_are_readable_private_and_do_not_issue(self):
        for doc in (self.quotation, self.request, self.report):
            for fmt in ('pdf', 'docx', 'xlsx'):
                with self.subTest(doctype=doc.doctype, format=fmt):
                    result = generate_document(doc.doctype, doc.name, output_format=fmt)
                    file = frappe.get_doc('File', result['file_id'])
                    self.assertEqual(file.is_private, 1)
                    self.assertEqual(file.attached_to_name, doc.name)
                    content = file.get_content()
                    if fmt == 'pdf': self.assertGreater(len(PdfReader(BytesIO(content)).pages), 0)
                    elif fmt == 'docx': self.assertTrue(Document(BytesIO(content)).tables)
                    else: self.assertTrue(load_workbook(BytesIO(content)).active.max_row)
        self.report.reload()
        self.assertEqual(self.report.status, 'Draft')
        self.assertFalse(self.report.pdf_file)

    def test_guest_and_unprivileged_user_cannot_generate_any_source(self):
        user = frappe.get_doc(dict(doctype='User', email=f'output-{uuid.uuid4().hex}@example.test',
            first_name='Output denied', send_welcome_email=0, roles=[dict(role='Employee')])).insert()
        before = frappe.db.count('File')
        for actor in ('Guest', user.name):
            frappe.set_user(actor)
            for doc in (self.quotation, self.request, self.report):
                for fmt in ('pdf', 'docx', 'xlsx'):
                    with self.assertRaises(frappe.PermissionError):
                        generate_document(doc.doctype, doc.name, output_format=fmt)
        frappe.set_user('Administrator')
        self.assertEqual(frappe.db.count('File'), before)

    def test_native_workspace_targets_exist_and_cards_are_preserved(self):
        workspace = frappe.get_doc('Workspace', 'Testing Center')
        self.assertEqual(len(workspace.number_cards), 6)
        for row in list(workspace.shortcuts) + list(workspace.links):
            if row.get('link_to'):
                self.assertTrue(frappe.db.exists('DocType', row.link_to), row.link_to)
        board = frappe.get_doc('Kanban Board', '检测主流程')
        self.assertEqual(len(board.columns), 8)
