"""Document output contracts without requiring a running Frappe site."""
import importlib
from copy import copy
from io import BytesIO
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


class OutputTests(unittest.TestCase):
    def test_registered_sources_and_default_pdf(self):
        from reli_lims.document_output.registry import get_template
        for dt in ('Quotation', 'Test Request', 'Test Report'):
            spec = get_template(dt)
            self.assertEqual(set(spec.formats), {'pdf', 'docx', 'xlsx'})
            for ext in ('html', 'docx', 'xlsx'):
                self.assertTrue((spec.directory / f'{spec.key}.{ext}').is_file())
        with self.assertRaises(ValueError): get_template('Customer')
        with self.assertRaises(ValueError): get_template('Quotation', '../../secret')

    def test_quotation_mapping_uses_native_totals_and_empty_fields(self):
        from reli_lims.document_output.mapping.quotation import build_context
        ctx = build_context(dict(name='Q-中文', customer_name='示例客户', currency='CNY',
            grand_total=99, items=[dict(item_name='振动', qty=2, rate=60, amount=90)]))
        self.assertEqual(ctx['rows'][0][-1], '90')
        self.assertIn(['价税合计', '99 CNY'], ctx['fields'])
        self.assertNotIn('None', str(ctx))

    def test_request_and_report_mapping_keep_samples_and_results(self):
        from reli_lims.document_output.mapping.commission import build_context
        from reli_lims.document_output.mapping.test_report import build_context as report_context
        ctx = build_context(dict(name='WT-1', samples=[dict(sample_code='S1', product_name='样品', quantity=2)],
            items=[dict(test_method='低温', standard='A', applicable_sample_codes='S1')]))
        self.assertIn('S1', str(ctx))
        self.assertIn('低温', str(ctx))
        report = report_context(dict(name='R-1', status='Draft'), dict(items=[
            dict(test_method='低温', standard='A', result='PASS', conclusion='符合要求')]))
        self.assertIn('草稿', report['notice'])
        self.assertIn('符合要求', str(report['rows']))
        self.assertNotIn('notes', report)  # internal notes are not customer deliverables

    def test_quotation_xlsx_preserves_template_and_expands_footer(self):
        from reli_lims.document_output.registry import get_template
        from reli_lims.document_output.renderers.xlsx_renderer import render
        from reli_lims.document_output.mapping.quotation import build_context
        spec = get_template('Quotation')
        ctx = build_context(dict(name='Q-1', customer_name='=HYPERLINK("bad")', grand_total=100,
            currency='CNY', items=[dict(item_name=f'项目{i}', qty=1, rate=10, amount=10) for i in range(8)]))
        before = load_workbook(spec.directory / 'quotation.xlsx').active
        sheet = load_workbook(BytesIO(render(spec, ctx))).active
        self.assertEqual(sheet['B4'].value, '=HYPERLINK("bad")')
        self.assertEqual(sheet['B4'].data_type, 's')
        self.assertEqual(sheet['B22'].value, '项目7')
        self.assertEqual(sheet['F23'].value, 100)
        self.assertIn('A42:F42', {str(r) for r in sheet.merged_cells.ranges})
        self.assertEqual(copy(sheet['B15'].font), copy(before['B15'].font))
        self.assertEqual(copy(sheet['B22'].border), copy(before['B16'].border))
        self.assertTrue(str(sheet.print_area))
        self.assertEqual(sheet['A42'].value, before['A36'].value)
        self.assertNotIn('晟楠', str([c.value for row in sheet for c in row]))

    def test_all_office_templates_handle_empty_and_many_rows(self):
        from reli_lims.document_output.registry import get_template
        from reli_lims.document_output.renderers.xlsx_renderer import render as xlsx
        from reli_lims.document_output.renderers.docx_renderer import render as docx
        from reli_lims.document_output.mapping.commission import build_context
        for count in (0, 1, 25):
            ctx = build_context(dict(name='中文文件名', items=[dict(test_method=f'检测{i}') for i in range(count)]))
            for dt in ('Test Request', 'Test Report', 'Quotation'):
                spec = get_template(dt)
                # context contracts have a fixed column count per source
                data = dict(ctx, rows=[[f'检测{i}'] * len(spec.columns) for i in range(count)], columns=spec.columns)
                for ext, output in [('xlsx', xlsx(spec, data)), ('docx', docx(spec, data))]:
                    if ext == 'xlsx':
                        text = str([c.value for row in load_workbook(BytesIO(output)).active for c in row])
                    else:
                        d = Document(BytesIO(output))
                        text = str([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
                    self.assertNotIn('{{', text)
                    if count: self.assertIn(f'检测{count-1}', text)

    def test_customer_literal_braces_survive_xlsx_template_replacement(self):
        from reli_lims.document_output.registry import get_template
        from reli_lims.document_output.renderers.xlsx_renderer import render
        from reli_lims.document_output.mapping.quotation import build_context
        literal = '客户约定 {{name}} 和 {{unknown}} 为样品代号'
        ctx = build_context(dict(name='Q', terms=literal))
        sheet = load_workbook(BytesIO(render(get_template('Quotation'), ctx))).active
        self.assertIn(literal, sheet['A20'].value)

    def test_long_xlsx_content_has_printable_continuation_and_wide_signatures(self):
        from reli_lims.document_output.registry import get_template
        from reli_lims.document_output.renderers.xlsx_renderer import render
        from reli_lims.document_output.mapping.quotation import build_context as quote
        from reli_lims.document_output.mapping.test_report import build_context as report
        long_text = '=HYPERLINK("untrusted")' + '合同条款内容。' * 800 + '全文结束标记'
        cases = [('Quotation', quote(dict(name='Q', terms=long_text))),
                 ('Test Report', report(dict(name='R', overall_conclusion=long_text), dict(items=[])))]
        for dt, context in cases:
            book = load_workbook(BytesIO(render(get_template(dt), context)))
            self.assertIn('续页', book.sheetnames)
            appendix = book['续页']
            self.assertIn(long_text, ''.join(str(c.value or '') for row in appendix for c in row))
            self.assertTrue(appendix.print_area)
            self.assertFalse(any(c.data_type == 'f' for row in appendix for c in row))
            self.assertTrue(all((dim.height or 0) < 409 for dim in appendix.row_dimensions.values()))
            self.assertIn('全文见续页', str([c.value for row in book.active for c in row]))
        sheet = load_workbook(BytesIO(render(get_template('Test Report'), report(dict(name='R'), dict(items=[]))))).active
        signature = next(c for row in sheet for c in row if c.value and '编制：' in str(c.value))
        self.assertIn(f'A{signature.row}:F{signature.row}', {str(r) for r in sheet.merged_cells.ranges})

    def test_pdf_escapes_markup_and_calls_native_pdf(self):
        from reli_lims.document_output.registry import get_template
        from reli_lims.document_output.renderers.pdf_renderer import render
        native = Mock(return_value=b'%PDF-test')
        modules = {'frappe': types.SimpleNamespace(), 'frappe.utils': types.ModuleType('frappe.utils'),
                   'frappe.utils.pdf': types.SimpleNamespace(get_pdf=native)}
        with patch.dict(sys.modules, modules):
            result = render(get_template('Test Request'), dict(title='<script>', name='WT', notice='',
                fields=[['客户', '<img src="http://secret">']], columns=['项目'], rows=[[0]], sections=[]))
        self.assertEqual(result, b'%PDF-test')
        html = native.call_args.args[0]
        self.assertNotIn('<script>', html)
        self.assertNotIn('<img', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertIn('<td>0</td>', html)


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.frappe = types.ModuleType('frappe')
        self.doc = Mock()
        self.doc.doctype = 'Test Request'
        self.doc.name = '中文/WT-1'
        self.doc.as_dict.return_value = dict(name=self.doc.name)
        self.frappe.get_doc = Mock(return_value=self.doc)
        self.files = types.ModuleType('frappe.utils.file_manager')
        self.files.save_file = Mock(return_value=types.SimpleNamespace(file_url='/private/files/output.pdf', name='FILE-1', file_name='output.pdf'))
        self.patcher = patch.dict(sys.modules, {'frappe': self.frappe, 'frappe.utils.file_manager': self.files})
        self.patcher.start()
        self.addCleanup(self.patcher.stop)
        sys.modules.pop('reli_lims.document_output.service', None)
        self.service = importlib.import_module('reli_lims.document_output.service')
        self.addCleanup(sys.modules.pop, 'reli_lims.document_output.service', None)

    def test_denied_source_cannot_map_render_or_write(self):
        self.doc.check_permission.side_effect = PermissionError('denied')
        with patch.object(self.service, 'render_document') as renderer:
            with self.assertRaises(PermissionError): self.service.generate_document('Test Request', 'WT')
            renderer.assert_not_called()
        self.files.save_file.assert_not_called()

    def test_private_attachment_pdf_default_and_filename(self):
        with patch.object(self.service, 'render_document', return_value=b'%PDF'):
            result = self.service.generate_document('Test Request', self.doc.name)
        self.assertEqual(result['output_format'], 'pdf')
        self.assertEqual(result['mime_type'], 'application/pdf')
        args, kwargs = self.files.save_file.call_args
        self.assertIn('中文', args[0])
        self.assertNotIn('/', args[0])
        self.assertEqual(kwargs['is_private'], 1)
        self.assertEqual(args[2:4], ('Test Request', self.doc.name))
        self.doc.check_permission.assert_any_call('read')
        self.doc.check_permission.assert_any_call('print')
        self.doc.apply_fieldlevel_read_permissions.assert_called_once()

    def test_link_labels_are_loaded_with_read_permission(self):
        self.doc.as_dict.return_value = dict(items=[dict(test_method='M-1', standard='S-1')])
        method = Mock()
        method.as_dict.return_value = dict(method_name='低温试验')
        standard = Mock()
        standard.as_dict.return_value = dict(full_name='GB/T 示例标准')
        self.frappe.get_doc.side_effect = [method, standard]
        ctx = self.service.build_context(self.doc)
        self.assertIn('低温试验', str(ctx['rows']))
        self.assertIn('GB/T 示例标准', str(ctx['rows']))
        method.check_permission.assert_called_with('read')
        standard.check_permission.assert_called_with('read')

    def test_invalid_format_template_and_source_do_not_load(self):
        for dt, template, fmt in [('Customer', None, 'pdf'), ('Test Request', '../../etc', 'pdf'),
                                  ('Test Request', None, 'zip')]:
            with self.assertRaises(ValueError): self.service.generate_document(dt, 'WT', template, fmt)
        self.frappe.get_doc.assert_not_called()

    def test_related_program_requires_permission(self):
        self.doc.doctype = 'Test Report'
        self.doc.as_dict.return_value = dict(test_program='P-1')
        program = Mock()
        program.check_permission.side_effect = PermissionError('program denied')
        self.frappe.get_doc.side_effect = [self.doc, program]
        with self.assertRaises(PermissionError): self.service.generate_document('Test Report', 'R-1')
        self.files.save_file.assert_not_called()
