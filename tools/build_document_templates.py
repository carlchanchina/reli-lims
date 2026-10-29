"""Rebuild shipped file templates using bundled Python dependencies.

Quotation starts from the supplied original workbook; pass --quotation-source.
Other templates are small editable A4 forms. Never use this on production data.
"""
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side

from reli_lims.document_output.registry import TEMPLATES

HTML = '''<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8"><style>
body{font-family:"Noto Sans CJK SC","WenQuanYi Micro Hei",sans-serif;font-size:11px;color:#111}
h1{font-size:22px;text-align:center}h2{font-size:14px;margin-top:18px}
p{white-space:pre-wrap}table{width:100%;border-collapse:collapse;table-layout:fixed;margin:12px 0}
th,td{border:1px solid #777;padding:6px;word-wrap:break-word;white-space:pre-wrap}
th{background:#f2f2f2}thead{display:table-header-group}tr{page-break-inside:avoid}
.notice{color:#555;text-align:center;font-size:10px}</style></head><body>
<h1>{{title}}</h1><p>单据编号：{{name}}</p><p class="notice">{{notice}}</p>
{{fields}}<h2>检测明细</h2>{{items}}{{sections}}
<p>编制：________________　审核：________________　日期：________________</p></body></html>'''


def docx_template(spec):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.5)
    style = doc.styles['Normal']
    style.font.name = 'Noto Sans CJK SC'
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    for key in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[key].font.color.rgb = RGBColor(0, 0, 0)
    # python-docx's bundled default can contain a blue Title border.
    for border in list(doc.styles['Title'].element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.add_paragraph('{{title}}', 'Title')
    doc.add_paragraph('单据编号：{{name}}')
    doc.add_paragraph('{{notice}}')
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = '{{field_label}}'
    table.rows[0].cells[1].text = '{{field_value}}'
    doc.add_heading('检测明细', level=2)
    table = doc.add_table(rows=2, cols=len(spec.columns))
    table.style = 'Table Grid'
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, column in enumerate(spec.columns):
        table.rows[0].cells[i].text = column
        table.rows[1].cells[i].text = '{{row.' + str(i) + '}}'
    if spec.key == 'quotation':
        doc.add_heading('报价条款', level=2)
        doc.add_paragraph('{{terms}}')
    section.footer.paragraphs[0].text = '编制：________________　审核：________________　日期：________________'
    doc.save(spec.directory / f'{spec.key}.docx')


def xlsx_template(spec):
    book = Workbook()
    sheet = book.active
    sheet.title = spec.title
    columns = len(spec.columns)
    end = chr(64 + columns)
    for r in (1, 3, 4, 5, 6, 7, 11):
        sheet.merge_cells(f'A{r}:{end}{r}')
    sheet['A1'] = '{{title}}'
    sheet['A3'] = '单据编号：{{name}}'
    sheet['A4'] = '{{notice}}'
    sheet['A6'] = '检测明细'
    for i, column in enumerate(spec.columns, 1):
        sheet.cell(8, i, column)
    sheet['A11'] = '单据详细信息'
    border = Border(*( [Side(style='thin', color='777777')] * 4 ))
    for row in sheet:
        for cell in row:
            cell.font = Font(name='宋体', size=11, bold=cell.row in (1, 8))
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if cell.row in (8, 9): cell.border = border
    sheet['A1'].font = Font(name='宋体', size=20, bold=True)
    for c in range(1, columns + 1): sheet.column_dimensions[chr(64+c)].width = 18
    for r in range(1, 12): sheet.row_dimensions[r].height = 30
    sheet.freeze_panes = 'B9'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '8:8'
    sheet.print_area = f'A1:{end}11'
    book.save(spec.directory / f'{spec.key}.xlsx')


def quotation_template(spec, source):
    book = load_workbook(source)
    sheet = book.active
    replacements = {'F3':'报价日期：{{date}}', 'B4':'{{customer}}', 'E4':'{{company}}',
        'B5':'{{contact}}', 'E5':'', 'B6':'{{address}}', 'E6':'{{company_address}}',
        'B7':'{{phone}}', 'E7':'', 'B8':'', 'E8':'', 'A9':'编号：{{name}}　币种：{{currency}}',
        'A11':'感谢您对本公司的信任和支持，检测报价如下。',
        'B14':'{{notice}}', 'A14':'状态', 'E17':'价税合计',
        'A19':'{{totals}}', 'A20':'报价条款：{{terms}}', 'A21':'', 'E23':'',
        'A27':'单位名称：{{company}}', 'A28':'', 'A29':'', 'A32':'{{company}}：'}
    for cell, value in replacements.items(): sheet[cell] = value
    for row in (15,16):
        for column in range(1,7): sheet.cell(row,column).value = None
    sheet['F17'] = '{{total}}'
    sheet['B13'] = '检测项目'
    sheet['E13'] = '数量'
    sheet['F13'] = '金额'
    # No print area was set in the original; explicitly cover its signature footer.
    sheet.print_area = 'A1:F37'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = '13:13'
    book.save(spec.directory / 'quotation.xlsx')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--quotation-source', type=Path, required=True)
    args = parser.parse_args()
    for spec in TEMPLATES.values():
        (spec.directory / f'{spec.key}.html').write_text(HTML)
        docx_template(spec)
        if spec.key == 'quotation': quotation_template(spec, args.quotation_source)
        else: xlsx_template(spec)


if __name__ == '__main__':
    main()
